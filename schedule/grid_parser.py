"""
Shared parser for the ВМК weekly-timetable "grid" layout, published both as
xlsx (e.g. groups' internal working copies) and as PDF (cs.msu.ru/studies/schedule).

Layout, in both formats:

    row: day name in col 1, group numbers in the remaining header columns
    row: time range in col 1 (spans 2 rows in xlsx / 1 merged cell in PDF),
         subject name + teacher/room text per group column

This module produces a common list-of-dict representation from either source,
so a single import pipeline can consume both. Cell text is kept verbatim in
`subject_raw` / `detail_raw` - `teachers` is a best-effort split (the source
documents are hand-typed and inconsistently formatted: teacher name + room
number + subgroup code all run together, separated by commas, newlines, or
nothing at all). Treat `teachers` as a convenience, `detail_raw` as ground
truth.

Not handled here: `mag_raspisanie_*.pdf` (non-integrated masters) - that file
is a different document entirely (no table grid, index/program layout) and
needs its own parser.
"""
import re

import openpyxl

WEEKDAYS = {"понедельник", "вторник", "среда", "четверг", "пятница", "суббота", "воскресенье"}

TITLE_PREFIX_RE = re.compile(
    r"^(профессор|доцент|академик\s+РАН|ст\.?\s*преп(?:одаватель)?\.?|"
    r"преподаватель|мл\.?\s*науч\.?\s*сотрудник|ассистент)\.?\s+",
    re.IGNORECASE,
)
ROOM_ONLY_RE = re.compile(r"^\d[\d\-–]*[а-яa-z]?$", re.IGNORECASE)
SUBGROUP_RE = re.compile(r"\b(МЗ-\d+)\b", re.IGNORECASE)
ROOM_RE = re.compile(r"\b(\d{2,4}[а-яa-z]?)\b")


def strip_title(name):
    """Drop a leading academic title ("профессор", "доцент", ...) from a name."""
    return TITLE_PREFIX_RE.sub("", name).strip()


def parse_detail(detail_raw):
    """Best-effort split of a free-text teacher/room/subgroup cell into entries."""
    if not detail_raw:
        return []
    fragments = [f.strip() for f in re.split(r"[,\n]+", detail_raw) if f.strip()]
    entries = []
    for frag in fragments:
        if ROOM_ONLY_RE.match(frag) and entries:
            entries[-1]["rooms"].append(frag)
            continue
        subgroups = SUBGROUP_RE.findall(frag)
        name_part = SUBGROUP_RE.sub("", frag)
        rooms = ROOM_RE.findall(name_part)
        for room in rooms:
            name_part = name_part.replace(room, "")
        name = strip_title(re.sub(r"\s{2,}", " ", name_part).strip(" .,"))
        if name or rooms or subgroups:
            entries.append({"name": name, "rooms": rooms, "subgroup": subgroups})
    return entries


def _split_time_range(time_range):
    parts = re.split(r"[–\-]", time_range, maxsplit=1)
    time_start = parts[0].strip() if parts else time_range.strip()
    time_end = parts[1].strip() if len(parts) > 1 else ""
    return time_start, time_end


def _cell_to_subject_detail(text):
    """Split a cell's combined "subject\\nteacher(s)" text into (subject, detail)."""
    if not text:
        return "", ""
    lines = [ln for ln in str(text).split("\n")]
    subject = lines[0].strip() if lines else ""
    detail = "\n".join(lines[1:]).strip() if len(lines) > 1 else ""
    return subject, detail


def _make_row(source, sheet, day, group, time_range, subject_raw, detail_raw):
    time_start, time_end = _split_time_range(time_range)
    return {
        "source": source,
        "sheet": sheet,
        "day": day,
        "group": group,
        "time_start": time_start,
        "time_end": time_end,
        "subject_raw": subject_raw,
        "detail_raw": detail_raw,
        "teachers": parse_detail(detail_raw),
    }


# --------------------------------------------------------------------------
# xlsx
# --------------------------------------------------------------------------

def _xlsx_merged_value(ws, row, col):
    cell = ws.cell(row=row, column=col)
    if cell.value is not None:
        return cell.value
    for merge_range in ws.merged_cells.ranges:
        if (row, col) in merge_range.cells:
            return ws.cell(row=merge_range.min_row, column=merge_range.min_col).value
    return None


def _xlsx_merge_span(ws, row, col):
    for merge_range in ws.merged_cells.ranges:
        if merge_range.min_col == col and merge_range.max_col == col and merge_range.min_row == row:
            return merge_range.max_row - merge_range.min_row + 1
    return 1


def parse_xlsx(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    rows_out = []
    seen = set()

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        max_row, max_col = ws.max_row, ws.max_column
        current_day = None
        group_cols = {}

        r = 1
        while r <= max_row:
            a_val = ws.cell(row=r, column=1).value
            if isinstance(a_val, str) and a_val.strip().lower() in WEEKDAYS:
                current_day = a_val.strip().lower()
                group_cols = {}
                for c in range(2, max_col + 1):
                    v = ws.cell(row=r, column=c).value
                    if isinstance(v, (int, float)):
                        group_cols[c] = int(v)
                r += 1
                continue

            if current_day and isinstance(a_val, str) and a_val.strip():
                time_range = a_val.strip()
                span = _xlsx_merge_span(ws, r, 1)
                detail_row = r + 1 if span >= 2 else r

                for c, group in group_cols.items():
                    subject = _xlsx_merged_value(ws, r, c)
                    detail = _xlsx_merged_value(ws, detail_row, c) if detail_row != r else None
                    if subject is None and detail is None:
                        continue
                    subject_raw = str(subject).strip() if subject else ""
                    detail_raw = str(detail).strip() if detail else ""
                    row_out = _make_row(path.name if hasattr(path, "name") else str(path),
                                        sheet_name, current_day, group, time_range,
                                        subject_raw, detail_raw)
                    key = (sheet_name, current_day, group, time_range, subject_raw, detail_raw)
                    if key in seen:
                        continue
                    seen.add(key)
                    rows_out.append(row_out)
                r += span
                continue

            r += 1

    return rows_out


# --------------------------------------------------------------------------
# pdf
# --------------------------------------------------------------------------

def parse_pdf(path):
    import pdfplumber

    rows_out = []
    settings = {
        "vertical_strategy": "lines",
        "horizontal_strategy": "lines",
        "text_x_tolerance": 2,
        "text_y_tolerance": 2,
    }

    with pdfplumber.open(path) as pdf:
        for page_index, page in enumerate(pdf.pages):
            table = page.extract_table(table_settings=settings)
            if not table:
                continue

            current_day = None
            group_cols = {}  # column index -> group label (e.g. "101", "541/1", "101бак")

            for raw_row in table:
                first = (raw_row[0] or "").strip() if raw_row[0] is not None else ""

                if first.lower() in WEEKDAYS:
                    current_day = first.lower()
                    group_cols = {}
                    for i, cell in enumerate(raw_row[1:], start=1):
                        if cell is None:
                            continue
                        label = str(cell).strip()
                        if not label or label.lower() in WEEKDAYS:
                            continue
                        if not re.match(r"^\d", label):
                            continue
                        # the source font occasionally kerns extra spaces inside digit
                        # runs ("1 11ио" for "111ио") - collapse those, but leave
                        # deliberate digit/letter separators alone
                        label = re.sub(r"(?<=\d)\s+(?=\d)", "", label)
                        group_cols[i] = label
                    continue

                if not current_day or not first:
                    continue

                time_range = first
                # forward-fill merged (None) cells; "" means "explicitly no lesson"
                filled = list(raw_row)
                last = None
                for i in range(1, len(filled)):
                    if filled[i] is None:
                        filled[i] = last
                    else:
                        last = filled[i]

                for i, group in group_cols.items():
                    cell_text = filled[i] if i < len(filled) else None
                    if not cell_text:
                        continue
                    subject_raw, detail_raw = _cell_to_subject_detail(cell_text)
                    if not subject_raw and not detail_raw:
                        continue
                    row_out = _make_row(path.name if hasattr(path, "name") else str(path),
                                        f"page{page_index + 1}", current_day, group,
                                        time_range, subject_raw, detail_raw)
                    rows_out.append(row_out)

    return rows_out


def parse_grid_file(path):
    """Dispatch by extension. `path` may be a str or Path."""
    from pathlib import Path
    p = Path(path)
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        return parse_xlsx(p)
    if p.suffix.lower() == ".pdf":
        return parse_pdf(p)
    raise ValueError(f"Unsupported file type: {p.suffix}")
